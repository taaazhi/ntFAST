"""
Автоматическое определение типа банка/источника по содержимому файла
Поддерживает: Kaspi, Halyk, Сбербанк, Binance и другие
"""
import logging
import re
import os
from typing import Optional, Dict, List, Tuple
from dataclasses import dataclass
from enum import Enum
import pdfplumber

logger = logging.getLogger(__name__)


class BankType(Enum):
    """Типы поддерживаемых банков и платформ"""
    KASPI = "kaspi"
    HALYK = "halyk"
    SBERBANK = "sberbank"
    JUSAN = "jusan"
    FORTE = "forte"
    BINANCE = "binance"
    FREEDOM = "freedom"
    UNKNOWN = "unknown"


@dataclass
class BankDetectionResult:
    """Результат определения банка"""
    bank_type: BankType
    confidence: float  # 0.0 - 1.0
    bank_name: str
    detected_keywords: List[str]
    metadata: Dict


class BankDetector:
    """
    Умный определитель типа банка по содержимому PDF
    Использует ключевые слова, структуру и паттерны
    """

    # Паттерны для определения банков
    BANK_PATTERNS = {
        BankType.KASPI: {
            "name": "Kaspi Bank",
            "keywords": [
                "kaspi", "каспи", "kaspi.kz", "kaspi gold", "kaspi bank",
                "АО «Kaspi Bank»", "Kaspi Pay", "KaspiPay",
                "Gold карта", "Kaspi Депозит"
            ],
            "account_patterns": [
                r"KZ\d{2}722C\d{12}",  # Kaspi IBAN
                r"\*\d{4}",  # Маскированная карта
            ],
            "structure_hints": [
                "Доступно на",
                "Пополнения",
                "Покупки",
                "Переводы",
                "Снятие наличных"
            ]
        },
        BankType.HALYK: {
            "name": "Halyk Bank",
            "keywords": [
                "halyk", "халык", "halykbank", "народный банк",
                "АО «Народный Банк Казахстана»", "Halyk Bank",
                "homebank", "хоум банк",
                "Народный Банк Казахстана",
                "Справка о наличии счета"
            ],
            "account_patterns": [
                r"KZ\d{2}601\d{13}",  # Halyk IBAN (KZ__601_____________)
            ],
            "structure_hints": [
                "Выписка по счету",
                "Народный Банк",
                "Справка о наличии счета",
                "Описаниеоперации",
                "Приходввалютесчета"
            ]
        },
        BankType.SBERBANK: {
            "name": "Сбербанк",
            "keywords": [
                "сбербанк", "sberbank", "сбер", "ПАО Сбербанк",
                "СберБанк Казахстан", "сбербанк онлайн"
            ],
            "account_patterns": [
                r"KZ\d{2}914\d{13}",  # Sber KZ IBAN
            ],
            "structure_hints": [
                "Выписка по карте",
                "Операции по счету"
            ]
        },
        BankType.JUSAN: {
            "name": "Jusan Bank",
            "keywords": [
                "jusan", "джусан", "jusan bank", "АТФБанк",
                "First Heartland Jusan Bank"
            ],
            "account_patterns": [
                r"KZ\d{2}826\d{13}",
            ],
            "structure_hints": []
        },
        BankType.FORTE: {
            "name": "ForteBank",
            "keywords": [
                "forte", "форте", "fortebank", "АО «ForteBank»"
            ],
            "account_patterns": [
                r"KZ\d{2}141\d{13}",
            ],
            "structure_hints": []
        },
        BankType.BINANCE: {
            "name": "Binance",
            "keywords": [
                "binance", "бинанс", "binance.com", "BNB",
                "USDT", "spot trading", "P2P", "криптовалюта"
            ],
            "account_patterns": [],
            "structure_hints": [
                "Trading History",
                "Deposit/Withdraw",
                "Spot Order History"
            ]
        },
        BankType.FREEDOM: {
            "name": "Freedom Finance",
            "keywords": [
                "freedom", "фридом", "freedom finance",
                "АО «Freedom Finance»", "tradernet"
            ],
            "account_patterns": [],
            "structure_hints": [
                "Брокерский отчет",
                "Инвестиционный счет"
            ]
        }
    }

    def __init__(self, pdf_path: str):
        self.pdf_path = pdf_path
        self.text_content = ""
        self.tables = []
        self._extract_content()

    def _is_xlsx(self) -> bool:
        """Проверить, является ли файл Excel XLSX"""
        return self.pdf_path.lower().endswith(('.xlsx', '.xls'))

    def _extract_content(self) -> None:
        """Извлечь текст и таблицы из файла для анализа"""
        if self._is_xlsx():
            self._extract_xlsx_content()
        else:
            self._extract_pdf_content()

    def _extract_xlsx_content(self) -> None:
        """Извлечь текст из XLSX для анализа (первые 15 строк).

        CRITICAL: read_only=False. Binance exports carry their metadata in a
        sheet whose dimensions openpyxl cannot determine lazily — in read-only
        mode iter_rows() yields nothing at all, the detector sees an empty
        document, scores every bank at 0 and falls back to GenericParser, which
        then extracts 0 transactions from a file BinanceParser reads fine.
        BinanceParser already learned this the hard way (see parsers/binance.py).
        """
        try:
            import openpyxl
            wb = openpyxl.load_workbook(self.pdf_path, read_only=False, data_only=True)
            sheet = wb.worksheets[0]
            lines = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= 15:
                    break
                row_text = " ".join(str(cell or "").strip() for cell in row if cell)
                if row_text:
                    lines.append(row_text)
            self.text_content = "\n".join(lines)
            wb.close()
        except Exception as e:
            logger.warning(f"Error extracting XLSX content: {e}")

    def _extract_pdf_content(self) -> None:
        """Извлечь текст и таблицы из PDF для анализа"""
        try:
            with pdfplumber.open(self.pdf_path) as pdf:
                # Извлекаем текст из первых 3 страниц (достаточно для определения)
                for page in pdf.pages[:3]:
                    text = page.extract_text() or ""
                    self.text_content += text + "\n"

                    # Извлекаем таблицы
                    tables = page.extract_tables()
                    self.tables.extend(tables)
        except Exception as e:
            logger.warning(f"Error extracting PDF content: {e}")

    def detect(self) -> BankDetectionResult:
        """
        Определить тип банка с указанием уверенности
        """
        scores: Dict[BankType, Tuple[float, List[str]]] = {}
        text_lower = self.text_content.lower()

        for bank_type, patterns in self.BANK_PATTERNS.items():
            score = 0.0
            found_keywords = []

            # Проверка ключевых слов (основной вес)
            for keyword in patterns["keywords"]:
                if keyword.lower() in text_lower:
                    score += 0.15
                    found_keywords.append(keyword)

            # Проверка паттернов номеров счетов
            for pattern in patterns["account_patterns"]:
                if re.search(pattern, self.text_content):
                    score += 0.25
                    found_keywords.append(f"account:{pattern}")

            # Проверка структурных подсказок
            for hint in patterns["structure_hints"]:
                if hint.lower() in text_lower:
                    score += 0.1
                    found_keywords.append(hint)

            # Нормализуем счет до 1.0
            score = min(score, 1.0)
            scores[bank_type] = (score, found_keywords)

        # Находим банк с наивысшим счетом
        best_bank = max(scores.items(), key=lambda x: x[1][0])
        bank_type = best_bank[0]
        confidence = best_bank[1][0]
        found_keywords = best_bank[1][1]

        # Если уверенность слишком низкая - неизвестный банк
        if confidence < 0.2:
            bank_type = BankType.UNKNOWN
            bank_name = "Неизвестный источник"
        else:
            bank_name = self.BANK_PATTERNS[bank_type]["name"]

        return BankDetectionResult(
            bank_type=bank_type,
            confidence=confidence,
            bank_name=bank_name,
            detected_keywords=found_keywords,
            metadata={
                "pages_analyzed": min(3, len(self.text_content.split('\n'))),
                "text_length": len(self.text_content),
                "tables_found": len(self.tables),
                "all_scores": {bt.value: s[0] for bt, s in scores.items()}
            }
        )

    def get_text_preview(self, max_chars: int = 500) -> str:
        """Получить превью текста для отладки"""
        return self.text_content[:max_chars]
